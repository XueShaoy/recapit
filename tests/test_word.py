from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from test_artifacts_formatter import transcription

from recapit.errors import ArtifactError
from recapit.formatter import render_markdown
from recapit.models import RecordingSummary
from recapit.word import markdown_to_docx

SAMPLE = """# 中文会议

- 来源：`会议.m4a`
- 语言：`zh`

## AI 摘要

这是**重要**且*清楚*的摘要。

## 关键结论

- 保留中文标点，完整表达。

## 待办事项

未识别到明确待办事项。

## 完整转写

[00:00:09] 第一段内容。
"""


def test_markdown_to_docx_preserves_structure_styles_and_text(tmp_path: Path) -> None:
    target = tmp_path / "中文会议.docx"
    markdown_to_docx(SAMPLE, target, replace_existing=False)

    document = Document(target)
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "中文会议",
        "来源：会议.m4a",
        "语言：zh",
        "AI 摘要",
        "这是重要且清楚的摘要。",
        "关键结论",
        "保留中文标点，完整表达。",
        "待办事项",
        "未识别到明确待办事项。",
        "完整转写",
        "[00:00:09] 第一段内容。",
    ]
    assert [paragraph.style.name for paragraph in document.paragraphs] == [
        "Title",
        "List Bullet",
        "List Bullet",
        "Heading 1",
        "Normal",
        "Heading 1",
        "List Bullet",
        "Heading 1",
        "Normal",
        "Heading 1",
        "Normal",
    ]
    summary = document.paragraphs[4]
    assert any(run.text == "重要" and run.bold for run in summary.runs)
    assert any(run.text == "清楚" and run.italic for run in summary.runs)
    source = document.paragraphs[1]
    assert any(run.text == "会议.m4a" and run.font.name == "Consolas" for run in source.runs)


def test_word_document_uses_a4_margins_and_east_asia_fonts(tmp_path: Path) -> None:
    target = tmp_path / "styled.docx"
    markdown_to_docx(SAMPLE, target, replace_existing=False)
    document = Document(target)
    section = document.sections[0]
    assert round(section.page_width.mm) == 210
    assert round(section.page_height.mm) == 297
    assert round(section.left_margin.mm) == 24
    assert round(section.top_margin.mm) == 22
    normal_fonts = document.styles["Normal"].element.rPr.rFonts
    title_fonts = document.styles["Title"].element.rPr.rFonts
    assert normal_fonts.get(qn("w:eastAsia")) == "Arial Unicode MS"
    assert title_fonts.get(qn("w:eastAsia")) == "Arial Unicode MS"


@pytest.mark.parametrize(
    ("markdown", "token"),
    [
        ("| A | B |\n| - | - |\n| 1 | 2 |\n", "table_open"),
        ("<section>内容</section>\n", "html_block"),
        ("> 引用\n", "blockquote_open"),
    ],
)
def test_word_export_rejects_unsupported_blocks(tmp_path: Path, markdown: str, token: str) -> None:
    with pytest.raises(ArtifactError, match=token):
        markdown_to_docx(markdown, tmp_path / "unsupported.docx", replace_existing=False)
    assert not (tmp_path / "unsupported.docx").exists()


def test_word_export_is_atomic_and_preserves_existing_file_on_replace_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    target = tmp_path / "会议.docx"
    target.write_bytes(b"original")

    def fail_replace(_source: str, _target: Path) -> None:
        raise OSError("injected replace failure")

    monkeypatch.setattr("recapit.word.os.replace", fail_replace)
    with pytest.raises(ArtifactError, match="injected replace failure"):
        markdown_to_docx(SAMPLE, target, replace_existing=True)

    assert target.read_bytes() == b"original"
    assert list(tmp_path.glob(".会议.docx.*.tmp")) == []


def test_word_export_refuses_overwrite_and_leaves_no_temporary_file(tmp_path: Path) -> None:
    target = tmp_path / "会议.docx"
    target.write_bytes(b"original")
    with pytest.raises(ArtifactError, match="已存在"):
        markdown_to_docx(SAMPLE, target, replace_existing=False)
    assert target.read_bytes() == b"original"
    assert list(tmp_path.glob(".会议.docx.*.tmp")) == []


@pytest.mark.parametrize(
    ("timestamps", "expected_codes"),
    [("paragraph", 2), ("segment", 3), ("none", 0)],
)
def test_word_matches_existing_timestamp_modes_and_empty_action_text(
    tmp_path: Path, timestamps: str, expected_codes: int
) -> None:
    value = transcription().with_summary(
        RecordingSummary(summary="摘要，保留中文标点。", key_points=[], action_items=[]),
        provider="agent",
        model=None,
    )
    markdown = render_markdown(
        value,
        timestamps=timestamps,  # type: ignore[arg-type]
        pause_seconds=2,
        max_chars=240,
    )
    target = tmp_path / f"{timestamps}.docx"
    markdown_to_docx(markdown, target, replace_existing=False)
    text = "\n".join(paragraph.text for paragraph in Document(target).paragraphs)
    assert text.count("[") == expected_codes
    assert "未识别到明确待办事项。" in text
    assert "摘要，保留中文标点。" in text
