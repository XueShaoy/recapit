from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from markdown_it import MarkdownIt
from markdown_it.token import Token

from recapit.errors import ArtifactError

_BODY_EAST_ASIA_FONT = "Arial Unicode MS"
_HEADING_EAST_ASIA_FONT = "Arial Unicode MS"
_SUPPORTED_BLOCK_TOKENS = {
    "bullet_list_close",
    "bullet_list_open",
    "heading_close",
    "heading_open",
    "inline",
    "list_item_close",
    "list_item_open",
    "paragraph_close",
    "paragraph_open",
}
_SUPPORTED_INLINE_TOKENS = {
    "code_inline",
    "em_close",
    "em_open",
    "hardbreak",
    "softbreak",
    "strong_close",
    "strong_open",
    "text",
}


def markdown_to_docx(markdown: str, path: Path, *, replace_existing: bool) -> None:
    """Convert Recapit's controlled Markdown subset to an atomically saved DOCX."""
    if path.exists() and not replace_existing:
        raise ArtifactError(f"目标产物已存在: {path}")

    document = _build_document(markdown)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary: str | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as stream:
            temporary = stream.name
        document.save(temporary)
        os.replace(temporary, path)
    except Exception as exc:
        if temporary:
            Path(temporary).unlink(missing_ok=True)
        raise ArtifactError(f"无法生成 Word 产物 {path}: {exc}") from exc


def _build_document(markdown: str) -> DocumentObject:
    parser = MarkdownIt("commonmark").enable("table")
    tokens = parser.parse(markdown)
    unsupported = next(
        (token.type for token in tokens if token.type not in _SUPPORTED_BLOCK_TOKENS),
        None,
    )
    if unsupported is not None:
        raise ArtifactError(f"Word 导出不支持 Markdown 块级结构: {unsupported}")

    document = Document()
    _configure_document(document)
    list_depth = 0
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type == "bullet_list_open":
            list_depth += 1
        elif token.type == "bullet_list_close":
            list_depth -= 1
        elif token.type in {"heading_open", "paragraph_open"}:
            if index + 2 >= len(tokens) or tokens[index + 1].type != "inline":
                raise ArtifactError(f"Word 导出遇到无效 Markdown token: {token.type}")
            inline = tokens[index + 1]
            if token.type == "heading_open":
                if token.tag == "h1":
                    paragraph = document.add_paragraph(style="Title")
                elif token.tag == "h2":
                    paragraph = document.add_paragraph(style="Heading 1")
                else:
                    raise ArtifactError(f"Word 导出不支持 Markdown 标题层级: {token.tag}")
            else:
                style = "List Bullet" if list_depth else "Normal"
                paragraph = document.add_paragraph(style=style)
            _append_inline(paragraph, inline.children or [])
            index += 2
        index += 1
    return document


def _append_inline(paragraph: Any, tokens: list[Token]) -> None:
    bold = False
    italic = False
    for token in tokens:
        if token.type not in _SUPPORTED_INLINE_TOKENS:
            raise ArtifactError(f"Word 导出不支持 Markdown 行内结构: {token.type}")
        if token.type == "strong_open":
            bold = True
        elif token.type == "strong_close":
            bold = False
        elif token.type == "em_open":
            italic = True
        elif token.type == "em_close":
            italic = False
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif token.type in {"text", "code_inline"}:
            run = paragraph.add_run(token.content)
            run.bold = bold
            run.italic = italic
            if token.type == "code_inline":
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                run_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
                run_fonts.set(qn("w:eastAsia"), _BODY_EAST_ASIA_FONT)


def _configure_document(document: DocumentObject) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(22)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(24)
    section.right_margin = Mm(24)

    _configure_style(document, "Normal", size=10.5, east_asia=_BODY_EAST_ASIA_FONT)
    _configure_style(document, "Title", size=22, east_asia=_HEADING_EAST_ASIA_FONT)
    _configure_style(document, "Heading 1", size=15, east_asia=_HEADING_EAST_ASIA_FONT)
    _configure_style(document, "List Bullet", size=10.5, east_asia=_BODY_EAST_ASIA_FONT)

    document.styles["Normal"].paragraph_format.space_after = Pt(7)
    document.styles["Normal"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    document.styles["Title"].paragraph_format.space_after = Pt(16)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    document.styles["List Bullet"].paragraph_format.space_after = Pt(3)


def _configure_style(document: DocumentObject, name: str, *, size: float, east_asia: str) -> None:
    style = document.styles[name]
    style.font.name = east_asia
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), east_asia)
    if name == "Title":
        paragraph_properties = style.element.get_or_add_pPr()
        border = paragraph_properties.find(qn("w:pBdr"))
        if border is not None:
            paragraph_properties.remove(border)
